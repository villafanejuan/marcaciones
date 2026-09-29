# -*- coding: utf-8 -*-
"""Captura las hojas del Excel real (Rooming) para el manual.

Requiere: openpyxl, pywin32 (pip install pywin32) y Microsoft Excel instalado.
Ejecutar con:  python docs/capturar_excel.py
"""
from __future__ import annotations

import ctypes
import os
import sys
import time

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

try:
    ctypes.windll.shcore.SetProcessDpiAwareness(2)
except Exception:
    try:
        ctypes.windll.user32.SetProcessDPIAware()
    except Exception:
        pass

import win32gui  # noqa: E402
from PIL import ImageGrab  # noqa: E402

from demo import construir_datos_demo  # noqa: E402
from excel_exportador import exportar_rooming_excel  # noqa: E402
from motor import OpcionesMotor, distribuir  # noqa: E402

CARPETA = os.path.join(os.path.dirname(os.path.abspath(__file__)), "imagenes")
os.makedirs(CARPETA, exist_ok=True)
RUTA_XLSX = os.path.join(CARPETA, "tmp_rooming_manual.xlsx")
RUTA_XLSX_CLARO = os.path.join(CARPETA, "tmp_rooming_manual_claro.xlsx")


def crear_copia_clara():
    """Copia del rooming con fondo blanco en celdas de datos (legible en cualquier
    tema). No modifica el archivo real de exportación."""
    import openpyxl
    from openpyxl.styles import PatternFill
    wb = openpyxl.load_workbook(RUTA_XLSX)
    blanco = PatternFill("solid", fgColor="FFFFFF")
    for hoja in wb.worksheets:
        for fila in hoja.iter_rows():
            for celda in fila:
                if celda.row == 1:
                    continue
                relleno = celda.fill
                tiene_fill = (
                    relleno is not None
                    and relleno.patternType == "solid"
                    and relleno.fgColor is not None
                    and relleno.fgColor.rgb not in (None, "00000000")
                )
                if not tiene_fill:
                    celda.fill = blanco
    wb.save(RUTA_XLSX_CLARO)


def _hwnd_excel():
    # La ventana principal de Excel tiene la clase "XLMAIN"
    encontradas = []
    def _cb(hwnd, extra):
        if win32gui.GetClassName(hwnd) == "XLMAIN":
            encontradas.append(hwnd)
        return True
    win32gui.EnumWindows(_cb, None)
    return encontradas[-1] if encontradas else None


def _traer_al_frente(hwnd):
    try:
        win32gui.ShowWindow(hwnd, 9)  # SW_RESTORE
        win32gui.SetForegroundWindow(hwnd)
        win32gui.BringWindowToTop(hwnd)
    except Exception:
        pass


def _título(hwnd):
    try:
        return win32gui.GetWindowText(hwnd)
    except Exception:
        return ""


def capturar_hoja(excel, libro, hoja: str, zoom: int, nombre: str) -> None:
    hoja_activa = libro.Worksheets(hoja)
    hoja_activa.Activate()
    hoja_activa.Select()
    excel.ActiveWindow.Zoom = zoom
    excel.ActiveWindow.DisplayGridlines = True
    excel.ScreenUpdating = True
    excel.Visible = True
    try:
        excel.ActiveWindow.ScrollRow = 1
        excel.ActiveWindow.ScrollColumn = 1
    except Exception:
        pass
    for _ in range(40):
        time.sleep(0.02)
    hwnd = _hwnd_excel()
    if hwnd is None:
        hwnd = excel.Hwnd if hasattr(excel, "Hwnd") else None
    if hwnd:
        _traer_al_frente(hwnd)
        time.sleep(0.8)
    for _ in range(40):
        time.sleep(0.02)
    rect = win32gui.GetWindowRect(hwnd) if hwnd else None
    if rect:
        x1, y1, x2, y2 = rect
        ancho = x2 - x1
        alto = y2 - y1
        imagen = ImageGrab.grab(bbox=(x1, y1, x2, y2))
        # recortar un margen mínimo para quitar bordes de la ventana, dejando
        # visible la cinta de Excel (reconoce que es Excel) y la fila 1.
        imagen = imagen.crop((2, 8, ancho - 2, alto - 2))
    else:
        imagen = ImageGrab.grab()
    ruta = os.path.join(CARPETA, nombre)
    imagen.save(ruta)
    print(f"  OK {nombre} ({imagen.size}) titulo='{_título(hwnd) if hwnd else ''}'")


def generar():
    hotel, grupos = construir_datos_demo()
    reporte = distribuir(hotel, grupos, OpcionesMotor())
    exportar_rooming_excel(RUTA_XLSX, reporte, destino="Camboriú")

    import win32com.client
    crear_copia_clara()
    excel = win32com.client.DispatchEx("Excel.Application")
    excel.Visible = True
    excel.DisplayAlerts = False
    try:
        libro = excel.Workbooks.Open(RUTA_XLSX_CLARO)
        capturar_hoja(excel, libro, "Límites por noche", 130, "06_excel_rooming.png")
        capturar_hoja(excel, libro, "Resumen por Grupo", 110, "07_excel_resumen.png")
        libro.Close(False)
    finally:
        excel.Quit()
    for ruta in (RUTA_XLSX, RUTA_XLSX_CLARO):
        if os.path.exists(ruta):
            os.remove(ruta)
    print("Listo. Imagenes en", CARPETA)


if __name__ == "__main__":
    generar()