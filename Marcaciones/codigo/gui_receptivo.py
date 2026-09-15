# -*- coding: utf-8 -*-
"""Interfaz grafica del sistema de planillas - Camboriu 2026/2027.

Permite generar las planillas de los 3 tipos (Receptivo, Operativo y Word),
elegir archivos, previsualizar cada hoja y guardar. Si el archivo de salida
esta abierto en Excel, se propone guardar una copia.
"""

import os
import glob
import threading
import tkinter as tk
from tkinter import ttk, filedialog, messagebox

import openpyxl

import generador_receptivo as gen
import generador_operativo as op
import generador_word as word

BASE = gen.BASE
ENTRADAS = gen.ENTRADAS
SALIDAS = gen.SALIDAS
MAX_PREVIEW_ROWS = 500

# tipo -> (entrada por defecto, plantilla por defecto, salida por defecto, ext)
def _primer_docx():
    found = glob.glob(os.path.join(ENTRADAS, "*.docx"))
    return found[0] if found else os.path.join(ENTRADAS, "OPERACIONES.docx")

TIPOS = {
    "Receptivo": dict(
        entrada=gen.IN_XLSX, plantilla=gen.PLANTILLA, salida=gen.OUT_XLSX,
        ext="xlsx", lbl="Entrada (colegios + marcaciones):"),
    "Operativo": dict(
        entrada=op.ARCHIVO_ENTRADA, plantilla=op.PLANTILLA,
        salida=op.ARCHIVO_SALIDA,
        ext="xlsx", lbl="Entrada (Control Operativo - hoja Reservas):"),
    "Word": dict(
        entrada=_primer_docx(), plantilla=word.PLANTILLA,
        salida=os.path.join(SALIDAS, "SALIDA_WORD.xlsx"),
        ext="docx", lbl="Entrada (documento de operaciones .docx):"),
}


class ReceptivoApp(tk.Tk):
    def __init__(self):
        super().__init__()
        self.title("Sistema de Planillas - Camboriu 2026/2027")
        self.geometry("1280x780")
        self.minsize(980, 620)

        self.var_tipo = tk.StringVar(value="Receptivo")
        self.var_entrada = tk.StringVar()
        self.var_plantilla = tk.StringVar()
        self.var_salida = tk.StringVar()

        self.wb = None
        self.stats = None
        self.hoja_actual = tk.StringVar()

        self._build_ui()
        self._aplicar_tipo()

    # ------------------------------------------------------------------ UI
    def _build_ui(self):
        estilo = ttk.Style(self)
        try:
            estilo.theme_use("vista")
        except tk.TclError:
            pass

        # selector de tipo
        top = ttk.Frame(self)
        top.pack(fill="x", padx=8, pady=(8, 2))
        ttk.Label(top, text="Tipo de planilla:").pack(side="left")
        self.cb_tipo = ttk.Combobox(top, textvariable=self.var_tipo,
                                    values=list(TIPOS.keys()),
                                    state="readonly")
        self.cb_tipo.pack(side="left", padx=6)
        self.cb_tipo.bind("<<ComboboxSelected>>", lambda e: self._aplicar_tipo())

        # panel de archivos
        cfg = ttk.LabelFrame(self, text=" Archivos ")
        cfg.pack(fill="x", padx=8, pady=4)
        cfg.columnconfigure(1, weight=1)

        self.lbl_entrada = tk.StringVar()
        self.lbl_plantilla = tk.StringVar(value="Plantilla (ejemplo del formato):")
        self.lbl_salida = tk.StringVar(value="Salida (archivo generado):")

        def fila_campo(parent, label_var, var, comando):
            ttk.Label(parent, textvariable=label_var).grid(sticky="w")
            ttk.Entry(parent, textvariable=var).grid(sticky="ew")
            ttk.Button(parent, text="Examinar...", command=comando).grid()
            parent.grid_columnconfigure(1, weight=1)

        fila_campo(cfg, self.lbl_entrada, self.var_entrada,
                   lambda: self._elegir_entrada())
        fila_campo(cfg, self.lbl_plantilla, self.var_plantilla,
                   lambda: self._elegir(self.var_plantilla, "xlsx", "Plantilla"))
        fila_campo(cfg, self.lbl_salida, self.var_salida,
                   lambda: self._elegir(self.var_salida, "xlsx", "Salida",
                                        ask_save=True))

        # botones de accion
        btns = ttk.Frame(self)
        btns.pack(fill="x", padx=8, pady=4)
        ttk.Button(btns, text="Generar planillas", command=self.generar).pack(side="left", padx=(0, 6))
        ttk.Button(btns, text="Guardar archivo salida", command=self.guardar).pack(side="left", padx=6)
        ttk.Button(btns, text="Guardar y abrir en Excel", command=self.guardar_y_abrir).pack(side="left", padx=6)
        ttk.Button(btns, text="Salir", command=self.destroy).pack(side="right")

        self.estado = tk.StringVar(value="Listo. Elija el tipo y los archivos, luego 'Generar planillas'.")
        ttk.Label(self, textvariable=self.estado, anchor="w").pack(fill="x", padx=8)

        # vista previa
        prev = ttk.LabelFrame(self, text=" Vista previa ")
        prev.pack(fill="both", expand=True, padx=8, pady=4)
        top2 = ttk.Frame(prev)
        top2.pack(fill="x", padx=4, pady=4)
        ttk.Label(top2, text="Hoja:").pack(side="left")
        self.cb_hojas = ttk.Combobox(top2, textvariable=self.hoja_actual, state="readonly")
        self.cb_hojas.pack(side="left", padx=4)
        self.cb_hojas.bind("<<ComboboxSelected>>", lambda e: self.mostrar_hoja())
        self.info_filas = ttk.Label(top2, text="", foreground="#555")
        self.info_filas.pack(side="right")

        cont = ttk.Frame(prev)
        cont.pack(fill="both", expand=True, padx=4, pady=(0, 4))
        tree = ttk.Treeview(cont, show="headings")
        vsb = ttk.Scrollbar(cont, orient="vertical", command=tree.yview)
        hsb = ttk.Scrollbar(cont, orient="horizontal", command=tree.xview)
        tree.configure(yscrollcommand=vsb.set, xscrollcommand=hsb.set)
        tree.grid(row=0, column=0, sticky="nsew")
        vsb.grid(row=0, column=1, sticky="ns")
        hsb.grid(row=1, column=0, sticky="ew")
        cont.grid_rowconfigure(0, weight=1)
        cont.grid_columnconfigure(0, weight=1)
        self.tree = tree

        self.resumen = ttk.Label(self, text="", anchor="w", foreground="#006600")
        self.resumen.pack(fill="x", padx=8, pady=(0, 6))

    # -------------------------------------------------------- tipo / helpers
    def _aplicar_tipo(self):
        tipo = self.var_tipo.get()
        cfg = TIPOS[tipo]
        self.lbl_entrada.set(cfg["lbl"])
        self.var_entrada.set(cfg["entrada"])
        self.var_plantilla.set(cfg["plantilla"])
        self.var_salida.set(cfg["salida"])
        self._cargar_hojas_disponibles()

    def _ext_entrada(self):
        return TIPOS[self.var_tipo.get()]["ext"]

    def _elegir_entrada(self):
        ext = self._ext_entrada()
        ft = [("Docx Word", "*.docx")] if ext == "docx" else [("Excel", "*.xlsx")]
        r = filedialog.askopenfilename(
            initialdir=os.path.dirname(self.var_entrada.get()) or BASE,
            filetypes=ft, title="Archivo de entrada")
        if r:
            self.var_entrada.set(r)
            if ext == "xlsx":
                self._cargar_hojas_disponibles()

    def _elegir(self, var, ext, titulo, ask_save=False):
        base = var.get() or BASE
        inicial = os.path.dirname(base) if os.path.dirname(base) else BASE
        ft = [("Excel", "*.xlsx")]
        if ask_save:
            r = filedialog.asksaveasfilename(initialdir=inicial, defaultextension=".xlsx",
                                             filetypes=ft, title=titulo)
        else:
            r = filedialog.askopenfilename(initialdir=inicial, filetypes=ft, title=titulo)
        if r:
            var.set(r)

    def _cargar_hojas_disponibles(self):
        try:
            lista = self._hojas_existentes()
        except Exception:
            lista = []
        if lista:
            self.cb_hojas["values"] = lista
            self.hoja_actual.set(lista[0])
            self.mostrar_hoja()

    def _hojas_existentes(self):
        ruta = self.var_entrada.get()
        if os.path.exists(ruta) and ruta.lower().endswith(".xlsx"):
            return openpyxl.load_workbook(ruta, read_only=True).sheetnames
        return []

    # ------------------------------------------------------------------ logica
    def generar(self):
        entrada = self.var_entrada.get()
        plantilla = self.var_plantilla.get()
        for etiqueta, ruta in (("Entrada", entrada), ("Plantilla", plantilla)):
            if not os.path.exists(ruta):
                messagebox.showerror("Falta archivo",
                                     f"'{etiqueta}' no existe:\n{ruta}")
                return
        self.estado.set("Generando...")
        self._run_worker(self._tarea_generar)

    def _tarea_generar(self):
        tipo = self.var_tipo.get()
        entrada = self.var_entrada.get()
        plantilla = self.var_plantilla.get()
        if tipo == "Receptivo":
            return gen.generar(entrada, plantilla)
        if tipo == "Operativo":
            return op.generar(entrada, plantilla)
        wb = word.generar(entrada, plantilla)
        return wb, self._stats_word(wb)

    @staticmethod
    def _stats_word(wb):
        atracc = {}
        for sh in wb.sheetnames:
            if sh.startswith("PROGRAMACION") or sh in ("5 COMIDA", "COLEGIOS MAESTRO"):
                continue
            ws = wb[sh]
            n = 0
            for r in range(3, ws.max_row + 1):
                b = ws.cell(row=r, column=2).value
                c = ws.cell(row=r, column=3).value
                if b is not None and c:
                    n += 1
            if n:
                atracc[sh] = n
        ws = wb["5 COMIDA"]
        n5 = sum(1 for r in range(4, ws.max_row + 1)
                 if ws.cell(row=r, column=2).value is not None)
        return {"colegios": 1, "atracciones": atracc, "5comida": n5}

    def guardar(self):
        self._guardar(abrir=False)

    def guardar_y_abrir(self):
        self._guardar(abrir=True)

    def _guardar(self, abrir):
        if self.wb is None:
            messagebox.showinfo("Nada que guardar", "Primero presione 'Generar planillas'.")
            return
        ruta = self._guardar_ahora(self.var_salida.get())
        if not ruta:
            if messagebox.askyesno(
                    "Archivo abierto en Excel",
                    "El archivo de salida no se puede sobrescribir "
                    "(probablemente esta abierto en Excel).\n\n"
                    "¿Quiere guardar una copia con otro nombre?"):
                copia = filedialog.asksaveasfilename(
                    initialdir=BASE, defaultextension=".xlsx",
                    filetypes=[("Excel", "*.xlsx")], title="Guardar copia de las planillas")
                if copia:
                    ruta = self._guardar_ahora(copia)
        if not ruta:
            return
        self.estado.set(f"Guardado: {ruta}")
        if abrir:
            os.startfile(ruta)

    def _guardar_ahora(self, ruta):
        try:
            self.wb.save(ruta)
            return ruta
        except PermissionError:
            self.estado.set(f"Error: '{os.path.basename(ruta)}' esta abierto o protegido.")
            return None
        except Exception as e:
            messagebox.showerror("Error al guardar", str(e))
            return None

    # ---------------------------------------------------------------- worker
    def _run_worker(self, func):
        def wrapper():
            try:
                resultado = func()
            except Exception as e:
                self.after(0, lambda: self._error(e))
                return
            self.after(0, lambda: self._ok(resultado))

        t = threading.Thread(target=wrapper, daemon=True)
        t.start()

    def _ok(self, resultado):
        wb, stats = resultado
        self.wb = wb
        self.stats = stats
        self.cb_hojas["values"] = wb.sheetnames
        self.hoja_actual.set(wb.sheetnames[0])
        self.mostrar_hoja()
        colegios = stats.get("colegios") or stats.get("grupos") or 0
        atr = ", ".join(f"{a} ({n})" for a, n in sorted(stats["atracciones"].items()))
        self.resumen.config(
            text=f"Grupos: {colegios}   |   5 COMIDA: {stats['5comida']} eventos   |"
                 f"   Atracciones: {atr}")
        self.estado.set("Planillas generadas. Revise en la vista previa y luego 'Guardar'.")

    def _error(self, e):
        self.estado.set("Error: " + str(e))
        messagebox.showerror("Error", str(e))

    # ----------------------------------------------------------------- preview
    def mostrar_hoja(self):
        if self.wb is None:
            return
        nombre = self.hoja_actual.get()
        if not nombre:
            return
        ws = self.wb[nombre]
        self.tree.delete(*self.tree.get_children())
        for c in self.tree["columns"]:
            self.tree.heading(c, text="")

        # aplana celdas unidas para que su valor se vea en todo el grupo
        efect = {}
        for mr in ws.merged_cells.ranges:
            v = ws.cell(row=mr.min_row, column=mr.min_col).value
            if v is None:
                continue
            for r in range(mr.min_row, mr.max_row + 1):
                for c in range(mr.min_col, mr.max_col + 1):
                    efect[(r, c)] = v

        total_filas = ws.max_row
        mostrar = min(total_filas, MAX_PREVIEW_ROWS)
        cols = list("ABCDEFGHIJKLMNOP")
        cols = cols[:max(1, ws.max_column)]
        self.tree["columns"] = ["#"] + cols
        self.tree.column("#", width=46, minwidth=40, stretch=False)
        self.tree.heading("#", text="Fila")
        for letra in cols:
            self.tree.column(letra, width=110, minwidth=48)
            self.tree.heading(letra, text=letra)

        for r in range(1, mostrar + 1):
            fila = [r]
            for c in range(1, len(cols) + 1):
                if (r, c) in efect:
                    v = efect[(r, c)]
                else:
                    v = ws.cell(row=r, column=c).value
                fila.append("" if v is None else v)
            self.tree.insert("", "end", values=fila)

        info = f"{total_filas} filas"
        if mostrar < total_filas:
            info += f" (mostrando primeras {mostrar})"
        self.info_filas.config(text=info)


if __name__ == "__main__":
    app = ReceptivoApp()
    app.mainloop()