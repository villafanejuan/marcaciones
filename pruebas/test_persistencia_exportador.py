from __future__ import annotations

import os
import sys
import unittest
from datetime import date

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from fabrica import mk_grupo, mk_hotel, mk_pax  # noqa: E402
from exportador import exportar_csv, exportar_resumen_csv  # noqa: E402
from excel_exportador import (  # noqa: E402
    exportar_rooming_excel,
    exportar_rooming_hoteles,
)
from modelos import (  # noqa: E402
    ConfigGrupo,
    ErrorDistribucion,
    Pax,
    Sexo,
    TipoPax,
)
from motor import OpcionesMotor, distribuir  # noqa: E402
from persistencia import (  # noqa: E402
    cargar_area,
    cargar_datos,
    grupo_a_dict,
    guardar_datos,
    guardar_multi,
    hotel_a_dict,
)

RUTA_TEMPORAL = os.path.join(
    os.path.dirname(os.path.abspath(__file__)), "tmp_datos_prueba.json"
)


class TestPersistencia(unittest.TestCase):
    def tearDown(self):
        if os.path.exists(RUTA_TEMPORAL):
            os.remove(RUTA_TEMPORAL)

    def _datos_ejemplo(self):
        hotel = mk_hotel([1, 2, 3, 4])
        hotel.nombre = "Hotel Roundtrip"
        hab = hotel.habitaciones[1]
        hab.piso = "2"
        hab.disponible_desde = date(2026, 1, 1)
        hab.disponible_hasta = date(2026, 12, 31)
        hab.agregar_bloqueo(date(2026, 4, 1), date(2026, 4, 5), "Feria")

        config = ConfigGrupo(
            separar_sexos=True,
            estudiantes_por_habitacion=4,
            estudiantes_estricto=True,
            max_ninos_por_habitacion=2,
            bebes_no_cuentan_cupo=True,
        )
        grupo = mk_grupo(
            "Grupo X",
            pax=[
                Pax(nombre="Ana", sexo=Sexo.F, pareja_de="Beto", notas="VIP"),
                Pax(nombre="Beto", sexo=Sexo.M, pareja_de="Ana"),
                Pax(nombre="Nini", tipo=TipoPax.NINO, edad=6, sexo=Sexo.F, familia="FAM"),
                Pax(nombre="Mama", sexo=Sexo.F, familia="FAM"),
            ],
            config=config,
            dia_inicio=3,
            noches=5,
        )
        return hotel, [grupo]

    def test_guardar_y_cargar_roundtrip(self):
        hotel_original, grupos_originales = self._datos_ejemplo()
        guardar_datos(RUTA_TEMPORAL, hotel_original, grupos_originales)
        hotel_cargado, grupos_cargados = cargar_datos(RUTA_TEMPORAL)
        self.assertEqual(hotel_a_dict(hotel_original), hotel_a_dict(hotel_cargado))
        self.assertEqual(
            [grupo_a_dict(g) for g in grupos_originales],
            [grupo_a_dict(g) for g in grupos_cargados],
        )

    def test_json_corrupto_error_amigable(self):
        with open(RUTA_TEMPORAL, "w", encoding="utf-8") as archivo:
            archivo.write("{esto no es json")
        with self.assertRaises(ErrorDistribucion):
            cargar_datos(RUTA_TEMPORAL)

    def test_archivo_inexistente_error_amigable(self):
        ruta = RUTA_TEMPORAL + ".no-existe.json"
        with self.assertRaises(ErrorDistribucion):
            cargar_datos(ruta)

    def test_enum_invalido_detectado(self):
        import json
        hotel_original, grupos_originales = self._datos_ejemplo()
        guardar_datos(RUTA_TEMPORAL, hotel_original, grupos_originales)
        with open(RUTA_TEMPORAL, encoding="utf-8") as archivo:
            datos = json.load(archivo)
        datos["hoteles"][0]["grupos"][0]["pax"][0]["tipo"] = "MARCIANO"
        with open(RUTA_TEMPORAL, "w", encoding="utf-8") as archivo:
            json.dump(datos, archivo)
        with self.assertRaises(ErrorDistribucion) as contexto:
            cargar_datos(RUTA_TEMPORAL)
        self.assertIn("TipoPax", str(contexto.exception))


class TestExportador(unittest.TestCase):
    def setUp(self):
        self.ruta_asignaciones = os.path.join(
            os.path.dirname(os.path.abspath(__file__)), "tmp_asignaciones.csv"
        )
        self.ruta_resumen = os.path.join(
            os.path.dirname(os.path.abspath(__file__)), "tmp_resumen.csv"
        )

    def tearDown(self):
        for ruta in (self.ruta_asignaciones, self.ruta_resumen):
            if os.path.exists(ruta):
                os.remove(ruta)

    def test_exportacion_csv_consistente(self):
        hotel = mk_hotel([2, 2])
        grupo = mk_grupo("G-CSV", [
            mk_pax("Uno", sexo=Sexo.M),
            mk_pax("Dos", sexo=Sexo.F),
            mk_pax("Tres", sexo=Sexo.X),
        ], config=ConfigGrupo(separar_sexos=False))
        reporte = distribuir(hotel, [grupo], OpcionesMotor())
        filas = exportar_csv(reporte, self.ruta_asignaciones)
        exportar_resumen_csv(reporte, self.ruta_resumen)
        self.assertEqual(filas, 3)
        with open(self.ruta_asignaciones, encoding="utf-8-sig") as archivo:
            lineas = archivo.read().strip().splitlines()
        self.assertEqual(len(lineas), 4)
        self.assertTrue(lineas[0].startswith("Grupo"))
        with open(self.ruta_resumen, encoding="utf-8-sig") as archivo:
            resumen_lineas = archivo.read().strip().splitlines()
        self.assertEqual(len(resumen_lineas), 2)
        self.assertIn("SI" if reporte.todo_alojado else "NO", resumen_lineas[1])

    def test_exportacion_con_faltantes(self):
        hotel = mk_hotel([2])
        grupo = mk_grupo("G-Corto", [mk_pax(f"P{i}") for i in range(5)],
                         config=ConfigGrupo(separar_sexos=False))
        reporte = distribuir(hotel, [grupo], OpcionesMotor())
        filas = exportar_csv(reporte, self.ruta_asignaciones)
        self.assertEqual(filas, 5)
        with open(self.ruta_asignaciones, encoding="utf-8-sig") as archivo:
            contenido = archivo.read()
        self.assertIn("SIN HABITACION", contenido)


class TestRoomingExcel(unittest.TestCase):
    RUTA = os.path.join(os.path.dirname(os.path.abspath(__file__)), "tmp_rooming.xlsx")

    def tearDown(self):
        if os.path.exists(self.RUTA):
            os.remove(self.RUTA)

    def test_resumen_y_detalle_por_categoria(self):
        from openpyxl import load_workbook
        hotel = mk_hotel([1, 2, 3, 4])
        grupo = mk_grupo("G-Rooming", [
            mk_pax("Juan", tipo=TipoPax.ESTUDIANTE, sexo=Sexo.M),
            mk_pax("Pepe", tipo=TipoPax.ESTUDIANTE, sexo=Sexo.M),
            mk_pax("Ana", tipo=TipoPax.ESTUDIANTE, sexo=Sexo.F),
            mk_pax("Luz", tipo=TipoPax.ESTUDIANTE, sexo=Sexo.F),
            mk_pax("Gui", tipo=TipoPax.GUIA, sexo=Sexo.M),
        ], config=ConfigGrupo(separar_sexos=True, estudiantes_por_habitacion=2))
        reporte = distribuir(hotel, [grupo], OpcionesMotor())
        filas = exportar_rooming_excel(self.RUTA, reporte, destino="Camboriú",
                                       titulo_empresa="Tudo Bem Eventos")
        self.assertGreater(filas, 0)
        libro = load_workbook(self.RUTA)
        self.assertEqual(libro.sheetnames,
                         ["Resumen por Grupo", "Detalle por Categoría",
                          "Cambios", "Límites por noche"])

        resumen = libro["Resumen por Grupo"]
        headers = [c.value for c in resumen[1]]
        self.assertIn("File", headers)
        self.assertIn("Colegio", headers)
        self.assertIn("Sin habitación", headers)
        self.assertIn("Estado", headers)
        fila = [c.value for c in resumen[2]]
        self.assertEqual(fila[0], "1")
        self.assertEqual(fila[1], "G-Rooming")
        self.assertEqual(fila[6], 5)

        detalle = libro["Detalle por Categoría"]
        celdas = list(detalle.iter_rows(values_only=True))
        categorias = [fila[4] for fila in celdas[1:] if fila[1] == "G-Rooming"]
        self.assertIn("Estudiantes - Hombres", categorias)
        self.assertIn("Estudiantes - Mujeres", categorias)
        self.assertIn("Guías (mixto)", categorias)

    def test_excepcion_por_faltantes(self):
        from openpyxl import load_workbook
        hotel = mk_hotel([2])
        grupo = mk_grupo("G-Fal", [mk_pax(f"P{i}") for i in range(4)],
                         config=ConfigGrupo(separar_sexos=False))
        reporte = distribuir(hotel, [grupo], OpcionesMotor())
        self.assertFalse(reporte.todo_alojado)
        exportar_rooming_excel(self.RUTA, reporte)
        libro = load_workbook(self.RUTA)
        detalle = libro["Detalle por Categoría"]
        celdas = list(detalle.iter_rows(values_only=True))
        fila_con_faltante = [f for f in celdas[1:] if f[13]]
        self.assertTrue(fila_con_faltante)

    def test_archivo_invalido_error_amigable(self):
        from modelos import ErrorDistribucion
        with self.assertRaises(ErrorDistribucion):
            exportar_rooming_excel(os.path.join(os.path.dirname(self.RUTA),
                                                "no-existe", "salida.xlsx"), None)

    def test_rooming_por_hotel_con_limites(self):
        from openpyxl import load_workbook
        import tempfile
        h1 = mk_hotel([2, 2])
        h1.nombre = "Hotel Norte"
        g1 = mk_grupo("G-A", [mk_pax("P1"), mk_pax("P2")],
                      config=ConfigGrupo(separar_sexos=False))
        h2 = mk_hotel([1, 2])
        h2.nombre = "Hotel Sur"
        g2 = mk_grupo("G-B", [mk_pax("P3"), mk_pax("P4")],
                      config=ConfigGrupo(separar_sexos=False))
        r1 = distribuir(h1, [g1], OpcionesMotor())
        r2 = distribuir(h2, [g2], OpcionesMotor())
        with tempfile.TemporaryDirectory() as tmp:
            rutas = exportar_rooming_hoteles(
                tmp, [h1, h2], [[g1], [g2]], reportes=[r1, r2], destino="Camboriú")
            self.assertEqual(len(rutas), 2)
            por_nombre = {os.path.basename(r): r for r in rutas}
            self.assertIn("Rooming_Hotel Norte.xlsx", por_nombre)
            self.assertIn("Rooming_Hotel Sur.xlsx", por_nombre)
            libro = load_workbook(por_nombre["Rooming_Hotel Norte.xlsx"])
            self.assertIn("Límites por noche", libro.sheetnames)
            ws = libro["Límites por noche"]
            encabezados = [c.value for c in ws[1]]
            self.assertIn("Límite personas", encabezados)
            self.assertIn("Límite aptos", encabezados)
            self.assertIn("Disponibles", encabezados)
            fila = [c.value for c in ws[2]]
            self.assertEqual(fila[1], "Hotel Norte")
            for c in ws.iter_rows(min_row=2, max_row=ws.max_row, max_col=8,
                                  values_only=True):
                self.assertIsNotNone(c[0])
                self.assertEqual(c[1], "Hotel Norte")

    def test_espacio_trabajo_con_estado(self):
        import tempfile
        ruta_json = os.path.join(tempfile.gettempdir(), "tmp_ws_estado.json")
        try:
            h1 = mk_hotel([2, 3])
            h1.nombre = "Hotel Uno"
            g1 = mk_grupo("G-C", [mk_pax(f"P{i}") for i in range(4)],
                          config=ConfigGrupo(separar_sexos=False))
            guardar_multi(ruta_json, [h1], [[g1]],
                          estado={"indice_hotel": 0,
                                  "vista_desde": "01/02/2026",
                                  "vista_hasta": "10/02/2026"})
            hoteles, grupos, estado = cargar_area(ruta_json)
            self.assertEqual([h.nombre for h in hoteles], ["Hotel Uno"])
            self.assertEqual(len(grupos[0]), 1)
            self.assertEqual(estado["indice_hotel"], 0)
            self.assertEqual(estado["vista_desde"], "01/02/2026")
        finally:
            if os.path.exists(ruta_json):
                os.remove(ruta_json)

    def test_roundtrip_empresa_provisional_confirmado_y_sin_fechas(self):
        hotel = mk_hotel([2, 3])
        g1 = mk_grupo("Con fechas", [mk_pax("A1"), mk_pax("A2")],
                      config=ConfigGrupo(separar_sexos=False),
                      empresa="Empresa Alfa", provisional=True, confirmado=False)
        g2 = mk_grupo("Sin fecha", [mk_pax("B1")],
                      dia_inicio=None, config=ConfigGrupo(separar_sexos=False),
                      empresa="Empresa Beta", provisional=True)
        g3 = mk_grupo("Confirmado", [mk_pax("C1"), mk_pax("C2")],
                      config=ConfigGrupo(separar_sexos=False), confirmado=True)
        guardar_datos(RUTA_TEMPORAL, hotel, [g1, g2, g3])
        _h, grupos_cargados = cargar_datos(RUTA_TEMPORAL)
        self.assertEqual(
            [grupo_a_dict(g) for g in [g1, g2, g3]],
            [grupo_a_dict(g) for g in grupos_cargados],
        )
        self.assertFalse(grupos_cargados[1].tiene_fechas)
        self.assertIsNone(grupos_cargados[1].fecha_ingreso)
        self.assertIsNone(grupos_cargados[1].fecha_egreso)
        self.assertEqual(grupos_cargados[0].empresa, "Empresa Alfa")
        self.assertTrue(grupos_cargados[0].provisional)
        self.assertFalse(grupos_cargados[0].confirmado)
        self.assertTrue(grupos_cargados[2].confirmado)

    def test_rooming_marca_estimado_confirmado_en_resumen(self):
        from openpyxl import load_workbook
        import tempfile
        with tempfile.TemporaryDirectory() as tmp:
            hotel = mk_hotel([2, 2, 2])
            hotel.nombre = "Hotel Estado"
            grupos = [
                mk_grupo("Prov", [mk_pax("p1"), mk_pax("p2")], dia_inicio=1,
                         noches=2, config=ConfigGrupo(separar_sexos=False),
                         provisional=True),
                mk_grupo("Conf", [mk_pax("q1"), mk_pax("q2")], dia_inicio=1,
                         noches=2, config=ConfigGrupo(separar_sexos=False),
                         confirmado=True),
            ]
            reporte = distribuir(hotel, grupos)
            rutas = exportar_rooming_hoteles(tmp, [hotel], [grupos], [reporte])
            libro = load_workbook(rutas[0])
            resumen = libro["Resumen por Grupo"]
            celdas = " ".join(str(c) for fila in resumen.iter_rows(values_only=True)
                              for c in fila if c is not None)
            self.assertIn("ESTIMADO", celdas)
            self.assertIn("CONFIRMADO", celdas)


if __name__ == "__main__":
    unittest.main()
